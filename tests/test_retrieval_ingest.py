from unittest.mock import MagicMock, patch

import pytest

from app.config import load_config
from app.retrieval.ingest import file_id_of, ingest_document
from app.retrieval.parsers import ParseError, parse_document

TEXT_MD = """# 面试知识库

多态是面向对象的三大特性之一。

重载是编译期多态，重写是运行期多态。
"""


# --- parsers ---


def test_parse_markdown(tmp_path):
    p = tmp_path / "kb.md"
    p.write_text(TEXT_MD, encoding="utf-8")
    text = parse_document(p)
    assert "# 面试知识库" in text
    assert "运行期多态" in text


def test_parse_txt_gbk(tmp_path):
    p = tmp_path / "kb.txt"
    p.write_bytes("中文内容".encode("gbk"))
    assert "中文内容" in parse_document(p)


def test_parse_unsupported_extension(tmp_path):
    p = tmp_path / "kb.csv"
    p.write_text("a,b", encoding="utf-8")
    with pytest.raises(ParseError, match="不支持的文档类型"):
        parse_document(p)


def test_parse_docx(tmp_path):
    import docx

    p = tmp_path / "kb.docx"
    d = docx.Document()
    d.add_paragraph("第一个段落")
    d.add_paragraph("第二个段落")
    d.save(str(p))
    text = parse_document(p)
    assert "第一个段落" in text
    assert "第二个段落" in text


def test_parse_pdf_with_text_layer(tmp_path):
    with patch("pdfplumber.open") as mock_open:
        page = MagicMock()
        page.extract_text.return_value = "PDF 第一页内容"
        pdf = MagicMock()
        pdf.pages = [page]
        mock_open.return_value.__enter__.return_value = pdf
        p = tmp_path / "kb.pdf"
        p.write_bytes(b"%PDF-1.4 fake")
        text = parse_document(p)
        assert "PDF 第一页内容" in text


def test_parse_pdf_without_text_layer(tmp_path):
    with patch("pdfplumber.open") as mock_open:
        page = MagicMock()
        page.extract_text.return_value = ""
        pdf = MagicMock()
        pdf.pages = [page]
        mock_open.return_value.__enter__.return_value = pdf
        p = tmp_path / "kb.pdf"
        p.write_bytes(b"%PDF-1.4 fake")
        with pytest.raises(ParseError, match="无文字层"):
            parse_document(p)


# --- ingest ---


def _managers():
    qdrant = MagicMock()
    qdrant.collection = "kb_blocks"
    es = MagicMock()
    es.index = "kb_blocks"
    embedding = MagicMock()
    embedding.dims = 1024
    embedding.embed_documents = MagicMock(side_effect=lambda texts: [[0.1] * 4 for _ in texts])
    return qdrant, es, embedding


def test_ingest_writes_parents_and_children(tmp_path):
    p = tmp_path / "kb.md"
    p.write_text(TEXT_MD, encoding="utf-8")
    qdrant, es, embedding = _managers()
    cfg = load_config()

    result = ingest_document(p, qdrant, es, embedding, cfg)

    assert result.parent_count >= 1
    assert result.child_count >= 1
    # ES 父块写入：文档 id = parent_id，含 file_id/text
    es.get_client().index.assert_called()
    _, kwargs = es.get_client().index.call_args
    assert kwargs["id"] == kwargs["document"]["parent_id"]
    assert kwargs["document"]["file_id"] == result.file_id
    # Qdrant 子块写入：向量 + payload（parent_id/text/seq/file_id）
    qdrant.get_client().upsert.assert_called()
    _, kwargs = qdrant.get_client().upsert.call_args
    assert kwargs["collection_name"] == "kb_blocks"
    points = kwargs["points"]
    assert len(points) == result.child_count
    assert all(len(pt["vector"]) == 4 for pt in points)
    assert all(pt["payload"]["file_id"] == result.file_id for pt in points)
    assert all(pt["payload"]["parent_id"] for pt in points)
    # 嵌入调用：子块文本批量
    embed_texts = embedding.embed_documents.call_args.args[0]
    assert len(embed_texts) == result.child_count


def test_ingest_idempotent_deletes_old_first(tmp_path):
    p = tmp_path / "kb.md"
    p.write_text(TEXT_MD, encoding="utf-8")
    qdrant, es, embedding = _managers()
    cfg = load_config()

    ingest_document(p, qdrant, es, embedding, cfg)
    file_id = file_id_of(p)
    qdrant.get_client().delete.assert_called_once()
    es.get_client().delete_by_query.assert_called_once()

    # 断言删除用了 file_id 过滤
    _, kwargs = qdrant.get_client().delete.call_args
    selector = kwargs["points_selector"]
    assert selector.must[0].match.value == file_id
    _, kwargs = es.get_client().delete_by_query.call_args
    assert kwargs["query"]["term"]["file_id"] == file_id


def test_ingest_creates_collection_and_index(tmp_path):
    p = tmp_path / "kb.md"
    p.write_text(TEXT_MD, encoding="utf-8")
    qdrant, es, embedding = _managers()
    cfg = load_config()

    ingest_document(p, qdrant, es, embedding, cfg)

    qdrant.ensure_collection.assert_called_once_with(1024)
    es.ensure_index.assert_called_once()


def test_ingest_empty_text_raises(tmp_path):
    p = tmp_path / "kb.md"
    p.write_text("", encoding="utf-8")
    qdrant, es, embedding = _managers()
    with pytest.raises(ParseError, match="无可提取文本"):
        ingest_document(p, qdrant, es, embedding, load_config())


def test_file_id_stable_per_path(tmp_path):
    p1 = tmp_path / "kb.md"
    p2 = tmp_path / "kb2.md"
    p1.write_text("a", encoding="utf-8")
    p2.write_text("b", encoding="utf-8")
    assert file_id_of(p1) == file_id_of(p1)
    assert file_id_of(p1) != file_id_of(p2)
    assert len(file_id_of(p1)) == 16
