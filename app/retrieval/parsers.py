"""轻量文档解析：MD/TXT/DOCX/PDF 四类 → 纯文本（架构 v0.4 轻量渐进方案）。"""

from pathlib import Path

from loguru import logger

SUPPORTED_EXTENSIONS = {".md", ".txt", ".docx", ".pdf"}


class ParseError(Exception):
    """解析失败（含 PDF 无文字层）。"""


def parse_document(path: str | Path) -> str:
    """按扩展名分发解析，返回规范化纯文本（统一换行、去零宽字符）。"""
    p = Path(path)
    ext = p.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        supported = "/".join(sorted(SUPPORTED_EXTENSIONS))
        raise ParseError(f"不支持的文档类型：{ext}（支持 {supported}）")
    raw = _parse_by_ext(p, ext)
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u200b", "").replace("\ufeff", "")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = text.strip()
    if not text:
        raise ParseError(f"文档无可提取文本：{p.name}")
    logger.info("parsed {name}: {chars} chars", name=p.name, chars=len(text))
    return text


def _parse_by_ext(p: Path, ext: str) -> str:
    if ext in (".md", ".txt"):
        return _read_text(p)
    if ext == ".docx":
        return _parse_docx(p)
    return _parse_pdf(p)


def _read_text(p: Path) -> str:
    for encoding in ("utf-8", "gbk"):
        try:
            return p.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError(f"无法识别文本编码：{p.name}")


def _parse_docx(p: Path) -> str:
    import docx  # 惰性导入：仅 .docx 解析时加载

    document = docx.Document(str(p))
    parts = [para.text for para in document.paragraphs if para.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _parse_pdf(p: Path) -> str:
    import pdfplumber  # 惰性导入：仅 .pdf 解析时加载

    parts: list[str] = []
    with pdfplumber.open(str(p)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
    joined = "\n\n".join(parts)
    if not joined.strip():
        raise ParseError(f"PDF 无文字层（扫描件需先 OCR 或提供文字版）：{p.name}")
    return joined
